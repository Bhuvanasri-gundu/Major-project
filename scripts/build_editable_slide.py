"""
Builds an editable, high-font-size slide for Stage 4 & Stage 5 and inserts it
into the user's PowerPoint presentation right after Slide 10.
Preserves all original slides without deleting anything.
"""

import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Paths
original_pptx = Path(r"C:\Users\VICTUS\.gemini\antigravity\brain\9a2af67b-3b4c-43d3-9aba-36deb348816c\.user_uploaded\media_1790493944688.pptx")
out_pptx_1 = Path(r"E:\Major_Project_Implementation\updated_presentation.pptx")
out_pptx_2 = Path(r"E:\Major_Project_Implementation\mhealth_presentation_with_editable_slide.pptx")

prs = Presentation(original_pptx)
initial_slide_count = len(prs.slides)
print(f"[INFO] Initial slides in presentation: {initial_slide_count}")

# Blank layout
blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0]
slide = prs.slides.add_slide(blank_layout)

# Colors matching presentation theme
CLR_TERRACOTTA = RGBColor(181, 76, 52)    # #B54C34
CLR_GREEN = RGBColor(46, 125, 50)         # #2E7D32
CLR_DARK = RGBColor(40, 40, 37)           # #282825
CLR_MUTED = RGBColor(90, 90, 85)          # #5A5A55
CLR_CARD_BG = RGBColor(255, 255, 255)     # #FFFFFF
CLR_CARD_BORDER = RGBColor(220, 224, 230) # #DCE0E6
CLR_BLUE = RGBColor(30, 64, 175)          # #1E40AF
CLR_GOLD = RGBColor(180, 83, 9)           # #B45309

# 1. Top terracotta accent stripe
stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
stripe.fill.solid()
stripe.fill.fore_color.rgb = CLR_TERRACOTTA
stripe.line.fill.background()

# 2. Slide Title & Subtitle (Text Box)
title_box = slide.shapes.add_textbox(Inches(0.65), Inches(0.24), Inches(12.0), Inches(0.95))
tf = title_box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

p1 = tf.paragraphs[0]
p1.text = "BIOMEDICAL CONCEPT EXTRACTION & MeSH NORMALIZATION"
p1.font.size = Pt(23)
p1.font.bold = True
p1.font.color.rgb = CLR_DARK

p2 = tf.add_paragraph()
p2.text = "Bridging Discovered Topics to Standardized Clinical Entities for Research Gap Detection"
p2.font.size = Pt(13)
p2.font.italic = True
p2.font.color.rgb = CLR_MUTED

# =========================================================================
# LEFT CARD: STAGE 4 — BIOMEDICAL CONCEPT EXTRACTION
# =========================================================================
c1_left = Inches(0.65)
c1_top = Inches(1.22)
c1_width = Inches(5.85)
c1_height = Inches(5.95)

# Card Container
card1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left, c1_top, c1_width, c1_height)
card1.fill.solid()
card1.fill.fore_color.rgb = CLR_CARD_BG
card1.line.color.rgb = CLR_CARD_BORDER
card1.line.width = Pt(1.5)

# Left color accent bar on Card 1
accent1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left, c1_top, Inches(0.14), c1_height)
accent1.fill.solid()
accent1.fill.fore_color.rgb = CLR_TERRACOTTA
accent1.line.fill.background()

# Card 1 Title & Tagline Box
hdr1_box = slide.shapes.add_textbox(c1_left + Inches(0.28), c1_top + Inches(0.16), c1_width - Inches(0.45), Inches(0.70))
tf1 = hdr1_box.text_frame
tf1.word_wrap = True
tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

p = tf1.paragraphs[0]
p.text = "STAGE 4 — BIOMEDICAL CONCEPT EXTRACTION"
p.font.size = Pt(16.5)
p.font.bold = True
p.font.color.rgb = CLR_TERRACOTTA

p = tf1.add_paragraph()
p.text = "Discover important medical ideas present inside research topics"
p.font.size = Pt(12)
p.font.color.rgb = CLR_MUTED

# Extraction Sources Box (Light Orange Tint)
src_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(0.28), c1_top + Inches(0.92), c1_width - Inches(0.50), Inches(0.85))
src_box.fill.solid()
src_box.fill.fore_color.rgb = RGBColor(254, 242, 237)
src_box.line.color.rgb = RGBColor(255, 204, 188)
src_box.line.width = Pt(1)
tf_src = src_box.text_frame
tf_src.word_wrap = True
tf_src.vertical_anchor = MSO_ANCHOR.MIDDLE

p = tf_src.paragraphs[0]
p.text = "Extraction Sources across 12,642 Papers:"
p.font.size = Pt(11.5)
p.font.bold = True
p.font.color.rgb = RGBColor(191, 54, 12)

p = tf_src.add_paragraph()
p.text = "• MeSH Descriptors   • Author Keywords   • Medical Noun Phrases"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = CLR_DARK

# Filtering Stats Container (3-Step Metric)
stat_box1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(0.28), c1_top + Inches(1.88), c1_width - Inches(0.50), Inches(1.30))
stat_box1.fill.solid()
stat_box1.fill.fore_color.rgb = RGBColor(248, 250, 252)
stat_box1.line.color.rgb = RGBColor(207, 216, 220)
stat_box1.line.width = Pt(1)

# Stat 1: 19,689
st1 = slide.shapes.add_textbox(c1_left + Inches(0.32), c1_top + Inches(1.98), Inches(1.70), Inches(1.10))
tf_st1 = st1.text_frame
tf_st1.word_wrap = True
p = tf_st1.paragraphs[0]
p.text = "19,689"
p.font.size = Pt(25)
p.font.bold = True
p.font.color.rgb = RGBColor(71, 85, 105)
p.alignment = PP_ALIGN.CENTER
p = tf_st1.add_paragraph()
p.text = "Raw Concepts"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = CLR_MUTED
p.alignment = PP_ALIGN.CENTER

# Arrow 1
arr1 = slide.shapes.add_textbox(c1_left + Inches(2.02), c1_top + Inches(2.15), Inches(0.40), Inches(0.50))
arr1.text_frame.paragraphs[0].text = "→"
arr1.text_frame.paragraphs[0].font.size = Pt(22)
arr1.text_frame.paragraphs[0].font.bold = True
arr1.text_frame.paragraphs[0].font.color.rgb = RGBColor(148, 163, 184)

# Stat 2: Frequency >= 5
st2 = slide.shapes.add_textbox(c1_left + Inches(2.32), c1_top + Inches(2.08), Inches(1.50), Inches(1.00))
tf_st2 = st2.text_frame
tf_st2.word_wrap = True
p = tf_st2.paragraphs[0]
p.text = "Freq ≥ 5"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = CLR_TERRACOTTA
p.alignment = PP_ALIGN.CENTER
p = tf_st2.add_paragraph()
p.text = "Occurrence Filter"
p.font.size = Pt(10.5)
p.font.bold = True
p.font.color.rgb = CLR_MUTED
p.alignment = PP_ALIGN.CENTER

# Arrow 2
arr2 = slide.shapes.add_textbox(c1_left + Inches(3.72), c1_top + Inches(2.15), Inches(0.40), Inches(0.50))
arr2.text_frame.paragraphs[0].text = "→"
arr2.text_frame.paragraphs[0].font.size = Pt(22)
arr2.text_frame.paragraphs[0].font.bold = True
arr2.text_frame.paragraphs[0].font.color.rgb = RGBColor(148, 163, 184)

# Stat 3: 2,790
st3 = slide.shapes.add_textbox(c1_left + Inches(4.02), c1_top + Inches(1.98), Inches(1.70), Inches(1.10))
tf_st3 = st3.text_frame
tf_st3.word_wrap = True
p = tf_st3.paragraphs[0]
p.text = "2,790"
p.font.size = Pt(25)
p.font.bold = True
p.font.color.rgb = CLR_GREEN
p.alignment = PP_ALIGN.CENTER
p = tf_st3.add_paragraph()
p.text = "Retained Concepts"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = CLR_GREEN
p.alignment = PP_ALIGN.CENTER

# Bullets Text Box for Card 1
b1_box = slide.shapes.add_textbox(c1_left + Inches(0.28), c1_top + Inches(3.28), c1_width - Inches(0.50), Inches(0.95))
tf_b1 = b1_box.text_frame
tf_b1.word_wrap = True
tf_b1.margin_left = tf_b1.margin_top = tf_b1.margin_right = tf_b1.margin_bottom = 0

p = tf_b1.paragraphs[0]
p.text = "• Topics identify broad research themes; concepts reveal exact medical ideas."
p.font.size = Pt(12.5)
p.font.color.rgb = CLR_DARK

p = tf_b1.add_paragraph()
p.text = "• Extracts specific clinical entities, technologies, symptoms, and therapies."
p.font.size = Pt(12.5)
p.font.color.rgb = CLR_DARK

# Easy Memory Callout Box (Yellow Background)
mem_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c1_left + Inches(0.28), c1_top + Inches(4.35), c1_width - Inches(0.50), Inches(1.40))
mem_box.fill.solid()
mem_box.fill.fore_color.rgb = RGBColor(254, 249, 231)
mem_box.line.color.rgb = RGBColor(245, 158, 11)
mem_box.line.width = Pt(1.5)
tf_mem = mem_box.text_frame
tf_mem.word_wrap = True
tf_mem.vertical_anchor = MSO_ANCHOR.MIDDLE

p = tf_mem.paragraphs[0]
p.text = "EASY MEMORY KEY:"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = CLR_GOLD

p = tf_mem.add_paragraph()
p.text = "★  Topics  =  Research Area   (e.g., Topic 1: HIV Prevention & Care)"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = CLR_DARK

p = tf_mem.add_paragraph()
p.text = "★  Concepts  =  Important Medical Ideas   (e.g., medication adherence)"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = CLR_DARK


# =========================================================================
# RIGHT CARD: STAGE 5 — MeSH NORMALIZATION
# =========================================================================
c2_left = Inches(6.83)
c2_top = Inches(1.22)
c2_width = Inches(5.85)
c2_height = Inches(5.95)

# Card Container
card2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left, c2_top, c2_width, c2_height)
card2.fill.solid()
card2.fill.fore_color.rgb = CLR_CARD_BG
card2.line.color.rgb = CLR_CARD_BORDER
card2.line.width = Pt(1.5)

# Left color accent bar on Card 2
accent2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left, c2_top, Inches(0.14), c2_height)
accent2.fill.solid()
accent2.fill.fore_color.rgb = CLR_GREEN
accent2.line.fill.background()

# Card 2 Title & Tagline Box
hdr2_box = slide.shapes.add_textbox(c2_left + Inches(0.28), c2_top + Inches(0.16), c2_width - Inches(0.45), Inches(0.70))
tf2 = hdr2_box.text_frame
tf2.word_wrap = True
tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0

p = tf2.paragraphs[0]
p.text = "STAGE 5 — MeSH NORMALIZATION"
p.font.size = Pt(16.5)
p.font.bold = True
p.font.color.rgb = CLR_GREEN

p = tf2.add_paragraph()
p.text = "Standardize concept names and eliminate duplicate lexical forms"
p.font.size = Pt(12)
p.font.color.rgb = CLR_MUTED

# Purpose Box (Light Green Tint)
purp_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(0.28), c2_top + Inches(0.92), c2_width - Inches(0.50), Inches(0.85))
purp_box.fill.solid()
purp_box.fill.fore_color.rgb = RGBColor(240, 253, 244)
purp_box.line.color.rgb = RGBColor(187, 247, 208)
purp_box.line.width = Pt(1)
tf_purp = purp_box.text_frame
tf_purp.word_wrap = True
tf_purp.vertical_anchor = MSO_ANCHOR.MIDDLE

p = tf_purp.paragraphs[0]
p.text = "Purpose: Ensure Consistency & Merge Synonyms"
p.font.size = Pt(11.5)
p.font.bold = True
p.font.color.rgb = RGBColor(22, 101, 52)

p = tf_purp.add_paragraph()
p.text = "• Merges plural forms, hyphens, and variants (e.g., 'cell phone' → 'smartphone')."
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = CLR_DARK

# Normalization Stats Container (3-Step Metric)
stat_box2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(0.28), c2_top + Inches(1.88), c2_width - Inches(0.50), Inches(1.30))
stat_box2.fill.solid()
stat_box2.fill.fore_color.rgb = RGBColor(248, 250, 252)
stat_box2.line.color.rgb = RGBColor(207, 216, 220)
stat_box2.line.width = Pt(1)

# Stat 1: 2,790
st4 = slide.shapes.add_textbox(c2_left + Inches(0.32), c2_top + Inches(1.98), Inches(1.65), Inches(1.10))
tf_st4 = st4.text_frame
tf_st4.word_wrap = True
p = tf_st4.paragraphs[0]
p.text = "2,790"
p.font.size = Pt(25)
p.font.bold = True
p.font.color.rgb = RGBColor(71, 85, 105)
p.alignment = PP_ALIGN.CENTER
p = tf_st4.add_paragraph()
p.text = "Extracted Concepts"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = CLR_MUTED
p.alignment = PP_ALIGN.CENTER

# Arrow 3
arr3 = slide.shapes.add_textbox(c2_left + Inches(1.92), c2_top + Inches(2.15), Inches(0.35), Inches(0.50))
arr3.text_frame.paragraphs[0].text = "→"
arr3.text_frame.paragraphs[0].font.size = Pt(22)
arr3.text_frame.paragraphs[0].font.bold = True
arr3.text_frame.paragraphs[0].font.color.rgb = RGBColor(148, 163, 184)

# Stat 2: 2,786
st5 = slide.shapes.add_textbox(c2_left + Inches(2.17), c2_top + Inches(1.98), Inches(1.70), Inches(1.10))
tf_st5 = st5.text_frame
tf_st5.word_wrap = True
p = tf_st5.paragraphs[0]
p.text = "2,786"
p.font.size = Pt(25)
p.font.bold = True
p.font.color.rgb = CLR_BLUE
p.alignment = PP_ALIGN.CENTER
p = tf_st5.add_paragraph()
p.text = "4 Duplicates Merged"
p.font.size = Pt(10.5)
p.font.bold = True
p.font.color.rgb = CLR_BLUE
p.alignment = PP_ALIGN.CENTER

# Arrow 4
arr4 = slide.shapes.add_textbox(c2_left + Inches(3.77), c2_top + Inches(2.15), Inches(0.35), Inches(0.50))
arr4.text_frame.paragraphs[0].text = "→"
arr4.text_frame.paragraphs[0].font.size = Pt(22)
arr4.text_frame.paragraphs[0].font.bold = True
arr4.text_frame.paragraphs[0].font.color.rgb = RGBColor(148, 163, 184)

# Stat 3: 21,439
st6 = slide.shapes.add_textbox(c2_left + Inches(4.02), c2_top + Inches(1.98), Inches(1.70), Inches(1.10))
tf_st6 = st6.text_frame
tf_st6.word_wrap = True
p = tf_st6.paragraphs[0]
p.text = "21,439"
p.font.size = Pt(25)
p.font.bold = True
p.font.color.rgb = CLR_GREEN
p.alignment = PP_ALIGN.CENTER
p = tf_st6.add_paragraph()
p.text = "Topic-Concept Pairs"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = CLR_GREEN
p.alignment = PP_ALIGN.CENTER

# Bullets Text Box for Card 2
b2_box = slide.shapes.add_textbox(c2_left + Inches(0.28), c2_top + Inches(3.28), c2_width - Inches(0.50), Inches(0.95))
tf_b2 = b2_box.text_frame
tf_b2.word_wrap = True
tf_b2.margin_left = tf_b2.margin_top = tf_b2.margin_right = tf_b2.margin_bottom = 0

p = tf_b2.paragraphs[0]
p.text = "• Establishes formal bipartite connections linking 63 topics to normalized MeSH terms."
p.font.size = Pt(12.5)
p.font.color.rgb = CLR_DARK

p = tf_b2.add_paragraph()
p.text = "• Avoids fragmented metrics by grouping morphological and synonym variations."
p.font.size = Pt(12.5)
p.font.color.rgb = CLR_DARK

# Pipeline Bridge Callout Box (Blue Background)
brg_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c2_left + Inches(0.28), c2_top + Inches(4.35), c2_width - Inches(0.50), Inches(1.40))
brg_box.fill.solid()
brg_box.fill.fore_color.rgb = RGBColor(239, 246, 255)
brg_box.line.color.rgb = RGBColor(59, 130, 246)
brg_box.line.width = Pt(1.5)
tf_brg = brg_box.text_frame
tf_brg.word_wrap = True
tf_brg.vertical_anchor = MSO_ANCHOR.MIDDLE

p = tf_brg.paragraphs[0]
p.text = "STRUCTURED PIPELINE BRIDGE:"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = CLR_BLUE

p = tf_brg.add_paragraph()
p.text = "Papers   ◄────────►   Topics   ◄────────►   Biomedical Concepts"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = RGBColor(29, 78, 216)

p = tf_brg.add_paragraph()
p.text = "Direct input for semantic relationship analysis and candidate research gap scoring."
p.font.size = Pt(11.5)
p.font.color.rgb = CLR_DARK

# Slide Number Text Box (bottom right)
num_box = slide.shapes.add_textbox(Inches(11.90), Inches(7.18), Inches(0.80), Inches(0.25))
p_num = num_box.text_frame.paragraphs[0]
p_num.text = "10b"
p_num.font.size = Pt(11)
p_num.font.color.rgb = CLR_MUTED

# Move the newly added slide to index 10 (Slide 11 in 1-based index)
slide_id_list = prs.slides._sldIdLst
new_slide_elem = slide_id_list[-1]
slide_id_list.insert(10, new_slide_elem)

# Save presentations
prs.save(out_pptx_1)
prs.save(out_pptx_2)
print(f"[SUCCESS] Saved updated presentations with {len(prs.slides)} slides (all original preserved).")
